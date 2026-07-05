"""Company onboarding from documents: brochure/price sheet in, company_ctx out.

    ctx = rsa.build_company_ctx("brochure.pdf", generate_fn=gen)   # LLM-structured block
    ctx = rsa.build_company_ctx("notes.txt")                       # no LLM: cleaned raw text

PDF needs pypdf, DOCX needs python-docx, XLSX needs openpyxl:
    pip install "rl-sales-augment[docs]"
"""
from __future__ import annotations
import os
import re

_PROMPT = (
    "Below is a company document. Write a compact COMPANY KNOWLEDGE block for a sales bot, "
    "plain text with these sections when the document supports them: Company (name, location, "
    "one-line story), Products (each with price if stated), Differentiators, Common objections "
    "and rebuttals, Anything a rep must never promise. Use ONLY facts from the document, keep "
    "every price and number exactly as written, no invention, no marketing fluff.\n\nDOCUMENT:\n"
)


def _load_text(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as e:
            raise ImportError('PDF support needs pypdf: pip install "rl-sales-augment[docs]"') from e
        return "\n".join((p.extract_text() or "") for p in PdfReader(path).pages)
    if ext in (".docx", ".doc"):
        try:
            import docx
        except ImportError as e:
            raise ImportError('DOCX support needs python-docx: pip install "rl-sales-augment[docs]"') from e
        return "\n".join(p.text for p in docx.Document(path).paragraphs)
    if ext in (".xlsx", ".xls"):
        try:
            import openpyxl
        except ImportError as e:
            raise ImportError('XLSX support needs openpyxl: pip install "rl-sales-augment[docs]"') from e
        wb = openpyxl.load_workbook(path, data_only=True)
        rows = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    rows.append(" | ".join(cells))
        return "\n".join(rows)
    with open(path, encoding="utf-8", errors="ignore") as fh:   # txt / md / csv
        return fh.read()


def build_company_ctx(path, generate_fn=None, max_chars=60000):
    """Document -> company_ctx string. With a generate_fn the LLM structures it; without,
    you get cleaned raw text (fine for short docs)."""
    raw = re.sub(r"\n{3,}", "\n\n", _load_text(path)).strip()[:max_chars]
    if not raw:
        raise ValueError(f"no text extracted from {path}")
    if generate_fn is None:
        return raw
    out = generate_fn(_PROMPT + raw)
    return (out or "").strip() or raw
